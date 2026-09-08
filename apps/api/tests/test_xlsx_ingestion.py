from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook
from openpyxl.packaging.relationship import Relationship
from openpyxl.workbook.external_link.external import ExternalBook, ExternalLink, ExternalSheetNames

from app.ingestion.errors import DatasetLimitError, EmptyDatasetError, MalformedDatasetError, SheetNotFoundError, UnsupportedFormatError
from app.ingestion.limits import IngestionLimits
from app.ingestion.service import ingest_dataset


def test_reads_valid_xlsx(xlsx_bytes) -> None:
    content = xlsx_bytes({"Customers": [["id", "name"], [1, "Ada"]]})
    dataset = ingest_dataset(content, "valid.xlsx")
    assert dataset.sheet_name == "Customers"
    assert dataset.row_count == 1
    assert dataset.column_names == ("id", "name")


def test_uses_first_visible_non_empty_sheet_and_allows_selection(xlsx_bytes) -> None:
    content = xlsx_bytes({"First": [["id"], [1]], "Second": [["name"], ["Ada"]]})
    assert ingest_dataset(content, "multiple_sheets.xlsx").sheet_name == "First"
    selected = ingest_dataset(content, "multiple_sheets.xlsx", sheet_name="Second")
    assert selected.column_names == ("name",)


def test_skips_hidden_first_sheet() -> None:
    workbook = Workbook()
    hidden = workbook.active
    hidden.title = "Hidden"
    hidden.append(["secret"])
    hidden.sheet_state = "hidden"
    visible = workbook.create_sheet("Visible")
    visible.append(["id"])
    visible.append([1])
    output = BytesIO()
    workbook.save(output)
    assert ingest_dataset(output.getvalue(), "hidden.xlsx").sheet_name == "Visible"


def test_rejects_missing_or_empty_sheet(xlsx_bytes) -> None:
    content = xlsx_bytes({"Only": [["id"], [1]]})
    with pytest.raises(SheetNotFoundError):
        ingest_dataset(content, "multiple_sheets.xlsx", sheet_name="Missing")
    empty = xlsx_bytes({"Empty": []})
    with pytest.raises(EmptyDatasetError):
        ingest_dataset(empty, "empty_sheet.xlsx")


def test_rejects_when_all_visible_sheets_are_empty(xlsx_bytes) -> None:
    with pytest.raises(EmptyDatasetError):
        ingest_dataset(xlsx_bytes({"One": [], "Two": []}), "all_empty.xlsx")


def test_formulas_are_not_evaluated(xlsx_bytes) -> None:
    content = xlsx_bytes({"Formula": [["id", "total", "name"], [1, "=1+1", "Ada"]]})
    dataset = ingest_dataset(content, "formulas.xlsx")
    assert dataset.ingestion_metadata.formula_cell_count == 1
    assert "formulas_not_evaluated" in dataset.warnings
    assert dataset.dataframe["total"][0] is None
    assert dataset.dataframe["name"][0] == "Ada"


def test_rejects_corrupt_workbook() -> None:
    with pytest.raises(MalformedDatasetError):
        ingest_dataset(b"not an xlsx", "corrupt.xlsx")


def test_enforces_sheet_and_cell_limits(xlsx_bytes) -> None:
    content = xlsx_bytes({"One": [["id"], [1]], "Two": [["id"], [2]]})
    with pytest.raises(DatasetLimitError):
        ingest_dataset(content, "sheets.xlsx", limits=IngestionLimits(max_sheets=1))
    cells = xlsx_bytes({"Data": [["a", "b"], [1, 2], [3, 4]]})
    with pytest.raises(DatasetLimitError):
        ingest_dataset(cells, "cells.xlsx", limits=IngestionLimits(max_total_cells=4))


def test_rejects_limit_before_default_sheet_scan(xlsx_bytes) -> None:
    oversized = xlsx_bytes({"Oversized": [["a", "b"], [1, 2]]})
    with pytest.raises(DatasetLimitError):
        ingest_dataset(oversized, "oversized.xlsx", limits=IngestionLimits(max_total_cells=3))


def test_enforces_zip_entry_and_expansion_limits() -> None:
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr("one.txt", "1")
        archive.writestr("two.txt", "2")
    with pytest.raises(DatasetLimitError):
        ingest_dataset(output.getvalue(), "entries.xlsx", limits=IngestionLimits(max_zip_entries=1))

    expanded = BytesIO()
    with ZipFile(expanded, "w", ZIP_DEFLATED) as archive:
        archive.writestr("large.txt", "x" * 100)
    with pytest.raises(DatasetLimitError):
        ingest_dataset(expanded.getvalue(), "expanded.xlsx", limits=IngestionLimits(max_zip_uncompressed_bytes=10))


def test_rejects_unsupported_office_formats() -> None:
    with pytest.raises(UnsupportedFormatError):
        ingest_dataset(b"anything", "legacy.xls")
    with pytest.raises(UnsupportedFormatError):
        ingest_dataset(b"anything", "macro.xlsm")


def test_rejects_malformed_internal_xml_as_a_safe_dataset_error(xlsx_bytes) -> None:
    source = xlsx_bytes({"Data": [["id"], [1]]})
    output = BytesIO()
    with ZipFile(BytesIO(source)) as source_archive, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for info in source_archive.infolist():
            payload = source_archive.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                payload = payload.replace(b"</worksheet>", b"<malformed")
            target.writestr(info, payload)

    with pytest.raises(MalformedDatasetError, match="malformed XML"):
        ingest_dataset(output.getvalue(), "malformed_xml.xlsx")

    missing_dimension = BytesIO()
    with ZipFile(BytesIO(source)) as source_archive, ZipFile(missing_dimension, "w", ZIP_DEFLATED) as target:
        for info in source_archive.infolist():
            payload = source_archive.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                start = payload.index(b"<dimension ")
                end = payload.index(b"/>", start) + 2
                payload = payload[:start] + payload[end:]
            target.writestr(info, payload)

    with pytest.raises(MalformedDatasetError, match="dimensions are missing"):
        ingest_dataset(missing_dimension.getvalue(), "missing_dimension.xlsx")


def test_external_workbook_relationship_is_not_loaded_or_exposed() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["id", "external_total"])
    worksheet.append([1, "='[SECRET_EXTERNAL.xlsx]Remote'!A1"])
    external_book = ExternalBook(
        sheetNames=ExternalSheetNames(sheetName=["Remote"]),
        id="rId1",
    )
    external_link = ExternalLink(externalBook=external_book)
    external_link.file_link = Relationship(
        Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/externalLinkPath",
        Target="https://example.invalid/SECRET_EXTERNAL.xlsx",
        TargetMode="External",
    )
    workbook._external_links.append(external_link)
    output = BytesIO()
    workbook.save(output)

    dataset = ingest_dataset(output.getvalue(), "external_link.xlsx")
    assert dataset.ingestion_metadata.formula_cell_count == 1
    assert dataset.dataframe["external_total"][0] is None
    assert "formulas_not_evaluated" in dataset.warnings
    assert "SECRET_EXTERNAL" not in repr(dataset)


def test_default_sheet_row_and_cell_limits_are_inclusive() -> None:
    workbook = Workbook()
    for index in range(20):
        worksheet = workbook.active if index == 0 else workbook.create_sheet()
        worksheet.title = f"Sheet{index}"
        worksheet.append(["id"])
        worksheet.append([index])
    output = BytesIO()
    workbook.save(output)
    assert ingest_dataset(output.getvalue(), "twenty_sheets.xlsx").sheet_name == "Sheet0"

    extra = workbook.create_sheet("Sheet20")
    extra.append(["id"])
    extra.append([20])
    output = BytesIO()
    workbook.save(output)
    with pytest.raises(DatasetLimitError, match="sheet limit"):
        ingest_dataset(output.getvalue(), "twenty_one_sheets.xlsx")

    rows = Workbook()
    worksheet = rows.active
    worksheet.append(["id"])
    worksheet.cell(row=100_001, column=1, value="boundary")
    output = BytesIO()
    rows.save(output)
    dataset = ingest_dataset(output.getvalue(), "one_hundred_thousand_rows.xlsx")
    assert dataset.row_count == 1

    worksheet.cell(row=100_002, column=1, value="over")
    output = BytesIO()
    rows.save(output)
    with pytest.raises(DatasetLimitError, match="row limit"):
        ingest_dataset(output.getvalue(), "over_one_hundred_thousand_rows.xlsx")

    cells = Workbook()
    worksheet = cells.active
    worksheet.append([f"c{index}" for index in range(250)])
    worksheet.cell(row=4_000, column=250, value="boundary")
    output = BytesIO()
    cells.save(output)
    dataset = ingest_dataset(output.getvalue(), "one_million_cells.xlsx")
    assert dataset.column_count == 250

    worksheet.cell(row=4_001, column=250, value="over")
    output = BytesIO()
    cells.save(output)
    with pytest.raises(DatasetLimitError, match="cell limit"):
        ingest_dataset(output.getvalue(), "over_one_million_cells.xlsx")


def test_default_zip_limits_are_inclusive(xlsx_bytes) -> None:
    source = xlsx_bytes({"Data": [["id"], [1]]})
    with ZipFile(BytesIO(source)) as archive:
        base_entry_count = len(archive.infolist())
        base_uncompressed_size = sum(info.file_size for info in archive.infolist())

    exact_entries = BytesIO()
    with ZipFile(BytesIO(source)) as source_archive, ZipFile(exact_entries, "w", ZIP_DEFLATED) as target:
        for info in source_archive.infolist():
            target.writestr(info, source_archive.read(info.filename))
        for index in range(2_000 - base_entry_count):
            target.writestr(f"security-padding/{index}", b"")
    assert ingest_dataset(exact_entries.getvalue(), "two_thousand_entries.xlsx").row_count == 1

    over_entries = BytesIO()
    with ZipFile(BytesIO(exact_entries.getvalue())) as source_archive, ZipFile(over_entries, "w", ZIP_DEFLATED) as target:
        for info in source_archive.infolist():
            target.writestr(info, source_archive.read(info.filename))
        target.writestr("security-padding/over", b"")
    with pytest.raises(DatasetLimitError, match="entry limit"):
        ingest_dataset(over_entries.getvalue(), "two_thousand_one_entries.xlsx")

    exact_expansion = BytesIO()
    padding_size = 20 * 1024 * 1024 - base_uncompressed_size
    with ZipFile(BytesIO(source)) as source_archive, ZipFile(exact_expansion, "w", ZIP_DEFLATED) as target:
        for info in source_archive.infolist():
            target.writestr(info, source_archive.read(info.filename))
        target.writestr("security-padding/exact", b"x" * padding_size)
    assert ingest_dataset(exact_expansion.getvalue(), "exact_expansion.xlsx").row_count == 1

    over_expansion = BytesIO()
    with ZipFile(BytesIO(source)) as source_archive, ZipFile(over_expansion, "w", ZIP_DEFLATED) as target:
        for info in source_archive.infolist():
            target.writestr(info, source_archive.read(info.filename))
        target.writestr("security-padding/over", b"x" * (padding_size + 1))
    with pytest.raises(DatasetLimitError, match="uncompressed size"):
        ingest_dataset(over_expansion.getvalue(), "over_expansion.xlsx")
