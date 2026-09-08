from io import BytesIO
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile, is_zipfile

import polars as pl
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .errors import DatasetLimitError, EmptyDatasetError, MalformedDatasetError, SheetNotFoundError
from .limits import IngestionLimits


def read_xlsx(
    content: bytes, limits: IngestionLimits, requested_sheet: str | None
) -> tuple[pl.DataFrame, tuple[str, ...], str, tuple[str, ...], int]:
    _validate_xlsx_archive(content, limits)
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except (BadZipFile, InvalidFileException, OSError, ParseError, ValueError) as error:
        raise MalformedDatasetError("XLSX workbook cannot be opened") from error
    try:
        if len(workbook.worksheets) > limits.max_sheets:
            raise DatasetLimitError("XLSX sheet limit exceeded")
        worksheet = _select_sheet(workbook, requested_sheet, limits)
        return _read_worksheet(worksheet, limits)
    except ParseError as error:
        raise MalformedDatasetError("XLSX workbook contains malformed XML") from error
    finally:
        workbook.close()


def _validate_xlsx_archive(content: bytes, limits: IngestionLimits) -> None:
    if not is_zipfile(BytesIO(content)):
        raise MalformedDatasetError("XLSX content is not a ZIP archive")
    try:
        with ZipFile(BytesIO(content)) as archive:
            infos = archive.infolist()
            if len(infos) > limits.max_zip_entries:
                raise DatasetLimitError("XLSX ZIP entry limit exceeded")
            if sum(info.file_size for info in infos) > limits.max_zip_uncompressed_bytes:
                raise DatasetLimitError("XLSX uncompressed size limit exceeded")
    except BadZipFile as error:
        raise MalformedDatasetError("XLSX archive is corrupt") from error


def _select_sheet(workbook: object, requested_sheet: str | None, limits: IngestionLimits) -> object:
    worksheets = workbook.worksheets  # type: ignore[attr-defined]
    if requested_sheet is not None:
        if requested_sheet not in workbook.sheetnames:  # type: ignore[attr-defined]
            raise SheetNotFoundError("Requested worksheet does not exist")
        worksheet = workbook[requested_sheet]  # type: ignore[index]
        _validate_worksheet_dimensions(worksheet, limits)
        return worksheet
    for worksheet in worksheets:
        if worksheet.sheet_state != "visible":
            continue
        _validate_worksheet_dimensions(worksheet, limits)
        if _has_values(worksheet):
            return worksheet
    raise EmptyDatasetError("Workbook has no visible non-empty worksheet")


def _validate_worksheet_dimensions(worksheet: object, limits: IngestionLimits) -> None:
    max_row = worksheet.max_row  # type: ignore[attr-defined]
    max_column = worksheet.max_column  # type: ignore[attr-defined]
    if max_row is None or max_column is None:
        raise MalformedDatasetError("XLSX worksheet dimensions are missing")
    if max_column > limits.max_columns:
        raise DatasetLimitError("XLSX column limit exceeded")
    if max_row > limits.max_rows + 1:
        raise DatasetLimitError("XLSX row limit exceeded")
    if max_row * max_column > limits.max_total_cells:
        raise DatasetLimitError("XLSX cell limit exceeded")


def _has_values(worksheet: object) -> bool:
    for row in worksheet.iter_rows(values_only=True):  # type: ignore[attr-defined]
        if any(value is not None for value in row):
            return True
    return False


def _read_worksheet(
    worksheet: object, limits: IngestionLimits
) -> tuple[pl.DataFrame, tuple[str, ...], str, tuple[str, ...], int]:
    _validate_worksheet_dimensions(worksheet, limits)
    rows: list[list[object | None]] = []
    formula_count = 0
    for row in worksheet.iter_rows():  # type: ignore[attr-defined]
        values: list[object | None] = []
        for cell in row:
            if cell.data_type == "f":
                formula_count += 1
                values.append(None)
            else:
                values.append(cell.value)
        rows.append(values)
    while rows and not any(value is not None for value in rows[-1]):
        rows.pop()
    if not rows or not any(value is not None for value in rows[0]):
        raise EmptyDatasetError("Worksheet has no header row")

    headers = tuple("" if value is None else str(value) for value in rows[0])
    _validate_headers(headers, limits)
    data_rows: list[list[object | None]] = []
    warnings: list[str] = []
    blank_rows_ignored = False
    for row in rows[1:]:
        if len(row) > len(headers) and any(value is not None for value in row[len(headers) :]):
            raise MalformedDatasetError("XLSX contains values beyond header columns")
        normalized = row[: len(headers)] + [None] * max(0, len(headers) - len(row))
        if not any(value is not None for value in normalized):
            blank_rows_ignored = True
            continue
        data_rows.append(normalized)
    if len(data_rows) > limits.max_rows or len(data_rows) * len(headers) > limits.max_total_cells:
        raise DatasetLimitError("XLSX dataset limit exceeded")
    if blank_rows_ignored:
        warnings.append("blank_xlsx_rows_ignored")
    if not data_rows:
        warnings.append("header_only_dataset")
    if formula_count:
        warnings.append("formulas_not_evaluated")
    dataframe = pl.DataFrame(data_rows, schema=list(headers), orient="row", strict=False)
    return dataframe, headers, worksheet.title, tuple(warnings), formula_count  # type: ignore[attr-defined]


def _validate_headers(headers: tuple[str, ...], limits: IngestionLimits) -> None:
    if not headers:
        raise EmptyDatasetError("Worksheet has no columns")
    if len(headers) > limits.max_columns:
        raise DatasetLimitError("XLSX column limit exceeded")
    if any(not header.strip() for header in headers):
        raise MalformedDatasetError("XLSX contains empty column names")
    if len(set(headers)) != len(headers):
        raise MalformedDatasetError("XLSX contains duplicate column names")
